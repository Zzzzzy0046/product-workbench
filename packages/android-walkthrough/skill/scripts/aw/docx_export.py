#!/usr/bin/env python3
"""Export product_summary.md to a readable DOCX with embedded screenshots."""

from __future__ import annotations

import argparse
import html
import re
from pathlib import Path
from urllib.parse import unquote

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

try:
    from PIL import Image
except Exception:  # pragma: no cover - script should still create text DOCX
    Image = None


IMG_RE = re.compile(r"<img\s+[^>]*src=[\"']([^\"']+)[\"'][^>]*>", re.I)
IMG_WIDTH_RE = re.compile(r"width=[\"']?(\d+)", re.I)
MD_IMG_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    for style_name, size, color in [
        ("Title", 22, "111827"),
        ("Heading 1", 18, "111827"),
        ("Heading 2", 15, "1D4ED8"),
        ("Heading 3", 12, "374151"),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)


def strip_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)
    tbl_pr.append(borders)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, margin: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "left", "bottom", "right"):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin))
        node.set(qn("w:type"), "dxa")


def resolve_path(raw: str, base_dir: Path) -> Path:
    value = unquote(raw.strip().strip('"').strip("'"))
    value = value.replace("/", "\\") if "\\" in value else value
    path = Path(value)
    if not path.is_absolute():
        path = base_dir / value
    return path.resolve()


def image_width_inches(path: Path, requested_px: int | None, max_width: float) -> float:
    if requested_px:
        return min(max(requested_px / 110.0, 1.0), max_width)
    if Image:
        try:
            with Image.open(path) as img:
                w, h = img.size
            if h > w:
                return min(1.55, max_width)
            return min(3.2, max_width)
        except Exception:
            pass
    return min(1.55, max_width)


def clean_inline(text: str) -> str:
    text = html.unescape(text)
    text = LINK_RE.sub(lambda m: f"{m.group(1)} ({m.group(2)})", text)
    return text


def add_formatted_runs(paragraph, text: str) -> None:
    text = clean_inline(text)
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=9.5, color="374151")
        else:
            run = paragraph.add_run(part)
            set_run_font(run)


def add_image_cell(cell, path: Path, width_in: float, caption: str) -> None:
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if path.exists():
        run = p.add_run()
        run.add_picture(str(path), width=Inches(width_in))
    else:
        run = p.add_run(f"[缺失图片: {path.name}]")
        set_run_font(run, size=9, color="B91C1C")


def add_image_group(doc: Document, image_specs: list[tuple[Path, int | None]], max_per_row: int) -> None:
    if not image_specs:
        return
    idx = 0
    while idx < len(image_specs):
        row_specs = image_specs[idx : idx + max_per_row]
        table = doc.add_table(rows=1, cols=len(row_specs))
        strip_table_borders(table)
        table.autofit = True
        max_width = 6.5 / max(1, len(row_specs))
        for cell, (path, requested_px) in zip(table.rows[0].cells, row_specs):
            width = image_width_inches(path, requested_px, max_width)
            add_image_cell(cell, path, width, path.stem)
        doc.add_paragraph()
        idx += max_per_row


def parse_img_specs(line: str, base_dir: Path) -> list[tuple[Path, int | None]]:
    specs = []
    for match in IMG_RE.finditer(line):
        full_tag = match.group(0)
        src = match.group(1)
        width_match = IMG_WIDTH_RE.search(full_tag)
        width_px = int(width_match.group(1)) if width_match else None
        specs.append((resolve_path(src, base_dir), width_px))
    return specs


def is_table_separator(line: str) -> bool:
    stripped = line.strip()
    if not stripped.startswith("|") or not stripped.endswith("|"):
        return False
    cells = [cell.strip() for cell in stripped.strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in cells)


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def normalize_table_header(value: str) -> str:
    return re.sub(r"[\s`*_：:]+", "", html.unescape(value)).lower()


def is_image_reference_cell(value: str) -> bool:
    text = re.sub(r"<[^>]+>", " ", value).strip()
    text = re.sub(r"[`*]+", "", text).strip()
    if not text:
        return True
    parts = [part.strip() for part in re.split(r"[\s,，、]+", text) if part.strip()]
    if not parts:
        return True
    return all(
        bool(re.search(r"\.(png|jpg|jpeg|webp|gif)$", part, re.I))
        or "/" in part
        or "\\" in part
        for part in parts
    )


def drop_docx_only_reference_columns(rows: list[list[str]]) -> list[list[str]]:
    if len(rows) < 2:
        return rows
    headers = rows[0]
    drop_indexes: set[int] = set()
    screenshot_headers = {"截图", "screenshot", "image", "图片", "证据图"}
    for idx, header in enumerate(headers):
        if normalize_table_header(header) not in screenshot_headers:
            continue
        cells = [(row[idx] if idx < len(row) else "") for row in rows[1:]]
        if all(is_image_reference_cell(cell) for cell in cells):
            drop_indexes.add(idx)

    if not drop_indexes:
        return rows
    return [[cell for idx, cell in enumerate(row) if idx not in drop_indexes] for row in rows]


def add_markdown_table(doc: Document, table_lines: list[str]) -> None:
    rows = [split_table_row(line) for line in table_lines if not is_table_separator(line)]
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    for row in rows:
        row.extend([""] * (col_count - len(row)))
    rows = drop_docx_only_reference_columns(rows)
    if not rows or not rows[0]:
        return
    col_count = max(len(row) for row in rows)
    for row in rows:
        row.extend([""] * (col_count - len(row)))

    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, "EAF2FF")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_formatted_runs(paragraph, value)
            for run in paragraph.runs:
                set_run_font(run, size=8.5 if col_count >= 5 else 9.5, bold=(r_idx == 0) or run.bold)
    doc.add_paragraph()


def section_heading(line: str) -> str | None:
    match = re.match(r"^##\s+(.+)$", line.strip())
    return match.group(1).strip() if match else None


def trim_html_image_block(block: list[str], max_images: int) -> list[str]:
    trimmed: list[str] = []
    kept = 0
    for line in block:
        if "<img" in line:
            if kept >= max_images:
                continue
            kept += 1
        trimmed.append(line)
    return trimmed


def extract_first_image_block(lines: list[str], max_images: int) -> list[str]:
    in_block = False
    block: list[str] = []
    for line in lines:
        if "<p" in line and not in_block:
            in_block = True
            block = [line]
            if "</p>" in line:
                return trim_html_image_block(block, max_images)
            continue
        if in_block:
            block.append(line)
            if "</p>" in line:
                return trim_html_image_block(block, max_images)
    return []


def extract_evaluation_bullets(lines: list[str], max_bullets: int) -> list[str]:
    labels = {"评价：", "评价:", "观察：", "观察:", "判断：", "判断:"}
    for idx, line in enumerate(lines):
        if line.strip() in labels:
            bullets: list[str] = []
            for later in lines[idx + 1 :]:
                stripped = later.strip()
                if stripped.startswith("- "):
                    bullets.append(later)
                    if len(bullets) >= max_bullets:
                        break
                elif bullets and stripped:
                    break
            if bullets:
                return bullets
    return [line for line in lines if line.strip().startswith("- ")][:max_bullets]


def extract_label_block(lines: list[str], labels: set[str], max_bullets: int = 8) -> list[str]:
    for idx, line in enumerate(lines):
        if line.strip() in labels:
            block = [line, ""]
            bullets = 0
            for later in lines[idx + 1 :]:
                stripped = later.strip()
                if stripped.startswith("- "):
                    block.append(later)
                    bullets += 1
                    if bullets >= max_bullets:
                        break
                elif not stripped:
                    if bullets:
                        break
                elif bullets:
                    break
            return block if bullets else []
    return []


def extract_subsection(lines: list[str], title: str) -> list[str]:
    start = next((idx for idx, line in enumerate(lines) if line.strip() == f"### {title}"), None)
    if start is None:
        return []
    end = len(lines)
    for idx in range(start + 1, len(lines)):
        if re.match(r"^#{2,4}\s+", lines[idx].strip()):
            end = idx
            break
    return lines[start:end]


def extract_current_conclusion(lines: list[str]) -> list[str]:
    result = []
    for line in lines:
        if line.strip().startswith("当前结论"):
            result.extend(["", line])
            break
    return result


def compact_monetization_section(heading: str, section_lines: list[str]) -> list[str]:
    compact: list[str] = [f"## {heading}", ""]
    compact.extend(extract_label_block(section_lines, {"观察到的触点：", "观察到的触点:"}, max_bullets=8))
    if compact[-2:] != ["", ""]:
        compact.append("")
    compact.extend(extract_label_block(section_lines, {"订阅页信息：", "订阅页信息:"}, max_bullets=8))

    for table_title in ("弹窗出现时机表", "广告出现时机表", "订阅出现时机表"):
        table = extract_subsection(section_lines, table_title)
        if table:
            if compact and compact[-1] != "":
                compact.append("")
            compact.extend(table)

    conclusion = extract_current_conclusion(section_lines)
    if conclusion:
        compact.extend(conclusion)

    while compact and not compact[-1].strip():
        compact.pop()
    compact.append("")
    return compact


def compact_markdown_lines(lines: list[str], max_images_per_section: int = 3) -> list[str]:
    first_section = next((idx for idx, line in enumerate(lines) if section_heading(line)), len(lines))
    compact: list[str] = []
    compact.extend(line for line in lines[:first_section] if not line.strip().startswith("- `"))
    while compact and not compact[-1].strip():
        compact.pop()
    compact.extend(["", "> Word 精简版：保留关键判断、核心模块评价和代表截图；完整证据链见 product_summary.md。", ""])

    sections: list[tuple[str, list[str]]] = []
    current_heading: str | None = None
    current_lines: list[str] = []
    for line in lines[first_section:]:
        heading = section_heading(line)
        if heading:
            if current_heading is not None:
                sections.append((current_heading, current_lines))
            current_heading = heading
            current_lines = [line]
        elif current_heading is not None:
            current_lines.append(line)
    if current_heading is not None:
        sections.append((current_heading, current_lines))

    monetization_headings = ("订阅", "评分", "广告", "商业化")
    for heading, section_lines in sections:
        if heading == "总体判断" or heading == "最终评价":
            compact.extend(section_lines)
            compact.append("")
            continue

        if any(keyword in heading for keyword in monetization_headings):
            compact.extend(compact_monetization_section(heading, section_lines))
            continue

        bullets = extract_evaluation_bullets(section_lines, max_bullets=3)
        images = extract_first_image_block(section_lines, max_images=max_images_per_section)
        if not bullets and not images:
            continue

        compact.append(f"## {heading}")
        compact.append("")
        compact.extend(bullets)
        if images:
            compact.extend(["", "代表截图：", ""])
            compact.extend(images)
        compact.append("")

    return compact


def export_docx(input_md: Path, output_docx: Path, max_per_row: int, mode: str, compact_images_per_section: int) -> tuple[int, int]:
    base_dir = input_md.parent
    lines = input_md.read_text(encoding="utf-8").splitlines()
    if mode == "compact":
        lines = compact_markdown_lines(lines, compact_images_per_section)

    doc = Document()
    configure_styles(doc)

    image_count = 0
    missing_count = 0
    pending_para: list[str] = []
    pending_table: list[str] = []
    in_html_image_block = False
    html_images: list[tuple[Path, int | None]] = []

    def flush_para() -> None:
        nonlocal pending_para
        text = " ".join(part.strip() for part in pending_para if part.strip()).strip()
        pending_para = []
        if text:
            p = doc.add_paragraph()
            add_formatted_runs(p, text)

    def flush_table() -> None:
        nonlocal pending_table
        if pending_table:
            add_markdown_table(doc, pending_table)
            pending_table = []

    def flush_html_images() -> None:
        nonlocal html_images, image_count, missing_count
        if html_images:
            image_count += len(html_images)
            missing_count += sum(1 for path, _ in html_images if not path.exists())
            add_image_group(doc, html_images, max_per_row)
        html_images = []

    for raw_line in lines:
        line = raw_line.rstrip()

        if pending_table and not line.strip().startswith("|"):
            flush_table()

        if "<p" in line and "<img" not in line:
            flush_table()
            flush_para()
            in_html_image_block = True
            html_images = []
            continue

        if in_html_image_block:
            html_images.extend(parse_img_specs(line, base_dir))
            if "</p>" in line:
                flush_html_images()
                in_html_image_block = False
            continue

        if "<img" in line:
            flush_table()
            flush_para()
            specs = parse_img_specs(line, base_dir)
            image_count += len(specs)
            missing_count += sum(1 for path, _ in specs if not path.exists())
            add_image_group(doc, specs, max_per_row)
            continue

        md_img_match = MD_IMG_RE.search(line)
        if md_img_match:
            flush_table()
            flush_para()
            path = resolve_path(md_img_match.group(2), base_dir)
            image_count += 1
            missing_count += 0 if path.exists() else 1
            add_image_group(doc, [(path, None)], max_per_row=1)
            continue

        if not line.strip():
            flush_table()
            flush_para()
            continue

        if line.strip().startswith("|") and line.strip().endswith("|"):
            flush_para()
            pending_table.append(line)
            continue

        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            flush_table()
            flush_para()
            level = len(heading.group(1))
            text = clean_inline(heading.group(2).strip())
            if level == 1:
                p = doc.add_paragraph(style="Title")
                add_formatted_runs(p, text)
            else:
                doc.add_heading(text, level=min(level - 1, 3))
            continue

        bullet = re.match(r"^[-*]\s+(.+)$", line)
        if bullet:
            flush_table()
            flush_para()
            p = doc.add_paragraph(style="List Bullet")
            add_formatted_runs(p, bullet.group(1).strip())
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if numbered:
            flush_table()
            flush_para()
            p = doc.add_paragraph(style="List Number")
            add_formatted_runs(p, numbered.group(1).strip())
            continue

        pending_para.append(line)

    flush_para()
    flush_table()
    flush_html_images()

    output_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_docx)
    return image_count, missing_count


def main() -> int:
    parser = argparse.ArgumentParser(description="Export product_summary.md to product_summary.docx.")
    parser.add_argument("--input-md", required=True, help="Path to product_summary.md")
    parser.add_argument("--out", help="Output .docx path. Defaults to product_summary.docx next to input.")
    parser.add_argument("--max-per-row", type=int, default=4, help="Max screenshots per image row.")
    parser.add_argument("--mode", choices=["compact", "full"], default="compact", help="compact keeps a review brief; full exports the entire markdown.")
    parser.add_argument("--compact-images-per-section", type=int, default=3, help="Representative screenshots per section in compact mode.")
    args = parser.parse_args()

    input_md = Path(args.input_md).resolve()
    if not input_md.exists():
        raise SystemExit(f"Missing input markdown: {input_md}")
    output = Path(args.out).resolve() if args.out else input_md.with_suffix(".docx")
    image_count, missing_count = export_docx(
        input_md,
        output,
        max(1, min(args.max_per_row, 6)),
        args.mode,
        max(0, min(args.compact_images_per_section, 8)),
    )
    print(output)
    print(f"mode={args.mode}")
    print(f"embedded_images={image_count}")
    print(f"missing_images={missing_count}")
    return 1 if missing_count else 0


if __name__ == "__main__":
    raise SystemExit(main())
