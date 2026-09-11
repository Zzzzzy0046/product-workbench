"""Small, deterministic exporters used by the local workbench.

Markdown remains the source of truth.  The exporters only render the selected
version and never ask the model to rewrite a document.
"""
from __future__ import annotations

import csv
import html
import io
import os
import re
import sys
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape as xml_escape


def _add_bundled_site_packages() -> None:
    """Use the bundled runtime packages when the workbench venv is minimal."""
    candidates = []
    configured = os.environ.get("WORKBENCH_EXPORT_PACKAGES", "").strip()
    if configured:
        candidates.append(Path(configured))
    candidates.append(
        Path.home()
        / ".cache"
        / "codex-runtimes"
        / "codex-primary-runtime"
        / "dependencies"
        / "python"
        / "Lib"
        / "site-packages"
    )
    for candidate in candidates:
        if candidate.is_dir() and str(candidate) not in sys.path:
            sys.path.insert(0, str(candidate))


def markdown_tables(markdown: str) -> list[tuple[str, list[list[str]]]]:
    """Return Markdown tables with the nearest preceding heading as title."""
    lines = str(markdown).splitlines()
    tables: list[tuple[str, list[list[str]]]] = []
    heading = "表格"
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        heading_match = re.match(r"^#{1,6}\s+(.+)$", line)
        if heading_match:
            heading = re.sub(r"[*_`]", "", heading_match.group(1)).strip() or heading
        if not line.startswith("|"):
            index += 1
            continue
        rows = []
        while index < len(lines) and lines[index].strip().startswith("|"):
            raw = lines[index].strip().strip("|")
            cells = [cell.strip() for cell in raw.split("|")]
            if not all(re.fullmatch(r"[:\-\s]+", cell or "-") for cell in cells):
                rows.append(cells)
            index += 1
        if rows:
            tables.append((heading, rows))
        continue
    return tables


def _asset_path(asset_resolver, source: str) -> Path | None:
    match = re.fullmatch(r"asset://([a-f0-9]{16,64})", str(source).strip())
    if not match:
        return None
    try:
        path = Path(asset_resolver(match.group(1)))
    except (OSError, TypeError, ValueError):
        return None
    return path if path.is_file() else None


def _plain_lines(markdown: str) -> list[tuple[str, str]]:
    """Yield a compact block representation for DOCX/PDF rendering."""
    result: list[tuple[str, str]] = []
    in_code = False
    for raw in str(markdown).splitlines():
        line = raw.rstrip()
        if line.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            result.append(("code", line))
            continue
        if not line.strip():
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            result.append((f"h{len(heading.group(1))}", re.sub(r"[*_`]", "", heading.group(2)).strip()))
        elif line.startswith("> "):
            result.append(("quote", line[2:]))
        elif re.match(r"^\s*[-*+]\s+", line):
            result.append(("bullet", re.sub(r"^\s*[-*+]\s+", "", line)))
        elif re.match(r"^\s*\d+[.)]\s+", line):
            result.append(("number", re.sub(r"^\s*\d+[.)]\s+", "", line)))
        elif line.startswith("|"):
            continue
        else:
            result.append(("p", line))
    return result


def _replace_inline(text: str) -> tuple[str, list[str]]:
    """Strip Markdown emphasis while returning asset references in order."""
    assets: list[str] = []

    def image(match: re.Match[str]) -> str:
        assets.append(match.group(2))
        return match.group(1) or "图片"

    text = re.sub(r"!\[([^\]]*)\]\((asset://[^)]+)\)", image, text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    text = re.sub(r"[*_`]+", "", text)
    return text, assets


def to_docx(markdown: str, asset_resolver=None) -> bytes:
    _add_bundled_site_packages()
    try:
        from docx import Document
        from docx.enum.text import WD_BREAK
        from docx.shared import Inches, Pt
    except ImportError as exc:  # pragma: no cover - depends on local runtime
        raise RuntimeError("Word 导出需要 python-docx；请安装文档导出依赖。") from exc

    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)
    for level, size in ((1, 18), (2, 15), (3, 12.5)):
        style = styles[f"Heading {level}"]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)

    lines = str(markdown).splitlines()
    index = 0
    in_code = False
    while index < len(lines):
        raw = lines[index].rstrip()
        if raw.startswith("```"):
            in_code = not in_code
            index += 1
            continue
        if in_code:
            paragraph = document.add_paragraph(style="No Spacing")
            run = paragraph.add_run(raw)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            index += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", raw)
        if heading:
            text, assets = _replace_inline(heading.group(2))
            p = document.add_heading(text, level=min(len(heading.group(1)), 3))
            for asset in assets:
                path = _asset_path(asset_resolver, asset) if asset_resolver else None
                if path:
                    p.add_run().add_break(WD_BREAK.LINE)
                    p.add_run().add_picture(str(path), width=Inches(5.8))
            index += 1
            continue
        if raw.strip().startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[:\-\s]+", cell or "-") for cell in cells):
                    rows.append(cells)
                index += 1
            if rows:
                table = document.add_table(rows=len(rows), cols=max(len(row) for row in rows))
                table.style = "Table Grid"
                for row_index, row in enumerate(rows):
                    for col_index, cell in enumerate(row):
                        table.cell(row_index, col_index).text = cell
            continue
        if not raw.strip():
            index += 1
            continue
        kind = "Normal"
        content = raw
        if raw.startswith("> "):
            kind, content = "Intense Quote", raw[2:]
        elif re.match(r"^\s*[-*+]\s+", raw):
            kind, content = "List Bullet", re.sub(r"^\s*[-*+]\s+", "", raw)
        elif re.match(r"^\s*\d+[.)]\s+", raw):
            kind, content = "List Number", re.sub(r"^\s*\d+[.)]\s+", "", raw)
        text, assets = _replace_inline(content)
        p = document.add_paragraph(style=kind if kind in styles else "Normal")
        p.add_run(text)
        for asset in assets:
            path = _asset_path(asset_resolver, asset) if asset_resolver else None
            if path:
                p.add_run().add_break(WD_BREAK.LINE)
                p.add_run().add_picture(str(path), width=Inches(5.8))
        index += 1
    result = io.BytesIO()
    document.save(result)
    return result.getvalue()


def to_pdf(markdown: str, asset_resolver=None) -> bytes:
    _add_bundled_site_packages()
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:  # pragma: no cover - depends on local runtime
        raise RuntimeError("PDF 导出需要 reportlab；请安装文档导出依赖。") from exc
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet()
    body = ParagraphStyle("WorkbenchBody", parent=styles["BodyText"], fontName="STSong-Light", fontSize=9.5, leading=15, alignment=TA_LEFT, spaceAfter=5)
    quote = ParagraphStyle("WorkbenchQuote", parent=body, leftIndent=10, textColor=colors.HexColor("#56625d"))
    heading_styles = {
        1: ParagraphStyle("H1", parent=body, fontSize=17, leading=22, spaceBefore=10, spaceAfter=8),
        2: ParagraphStyle("H2", parent=body, fontSize=13.5, leading=18, spaceBefore=8, spaceAfter=6),
        3: ParagraphStyle("H3", parent=body, fontSize=11.5, leading=16, spaceBefore=6, spaceAfter=4),
    }
    flow = []
    lines = str(markdown).splitlines()
    index = 0
    while index < len(lines):
        raw = lines[index].strip()
        if not raw:
            index += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", raw)
        if heading:
            text, assets = _replace_inline(heading.group(2))
            flow.append(Paragraph(html.escape(text), heading_styles.get(min(len(heading.group(1)), 3), heading_styles[3])))
            for asset in assets:
                path = _asset_path(asset_resolver, asset) if asset_resolver else None
                if path:
                    flow.append(Image(str(path), width=150 * mm, height=100 * mm, kind="proportional"))
            index += 1
            continue
        if raw.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[:\-\s]+", cell or "-") for cell in cells):
                    rows.append([Paragraph(html.escape(cell), body) for cell in cells])
                index += 1
            if rows:
                table = Table(rows, repeatRows=1, hAlign="LEFT")
                table.setStyle(TableStyle([
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#c8d0ca")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#edf2ed")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                ]))
                flow.extend([table, Spacer(1, 5)])
            continue
        if raw.startswith("> "):
            kind, content = "quote", raw[2:]
        elif re.match(r"^\s*[-*+]\s+", raw):
            kind, content = "bullet", "• " + re.sub(r"^\s*[-*+]\s+", "", raw)
        elif re.match(r"^\s*\d+[.)]\s+", raw):
            kind, content = "number", re.sub(r"^\s*\d+[.)]\s+", "", raw)
        else:
            kind, content = "p", raw
        text, assets = _replace_inline(content)
        flow.append(Paragraph(html.escape(text), quote if kind == "quote" else body))
        for asset in assets:
            path = _asset_path(asset_resolver, asset) if asset_resolver else None
            if path:
                flow.append(Image(str(path), width=150 * mm, height=100 * mm, kind="proportional"))
        index += 1
    result = io.BytesIO()
    SimpleDocTemplate(result, pagesize=A4, rightMargin=16 * mm, leftMargin=16 * mm, topMargin=15 * mm, bottomMargin=15 * mm, title="产品工作台交付物").build(flow)
    return result.getvalue()


def to_xlsx(markdown: str) -> bytes:
    _add_bundled_site_packages()
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - depends on local runtime
        raise RuntimeError("Excel 导出需要 openpyxl；请安装表格导出依赖。") from exc
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    tables = markdown_tables(markdown)
    if not tables:
        tables = [("正文", [["内容"], *[[line] for line in str(markdown).splitlines() if line.strip()]])]
    used_names: set[str] = set()
    for title, rows in tables:
        base = re.sub(r"[\\/*?:\[\]]", "", title).strip() or "表格"
        base = base[:31]
        name = base
        suffix = 2
        while name in used_names:
            name = f"{base[:28]}-{suffix}"
            suffix += 1
        used_names.add(name)
        sheet = workbook.create_sheet(name)
        for row_index, row in enumerate(rows, 1):
            for col_index, value in enumerate(row, 1):
                cell = sheet.cell(row=row_index, column=col_index, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_index == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="244C3E")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for col_index in range(1, max(len(row) for row in rows) + 1):
            width = min(48, max(12, max(len(str(row[col_index - 1])) if col_index <= len(row) else 0 for row in rows) + 2))
            sheet.column_dimensions[get_column_letter(col_index)].width = width
    result = io.BytesIO()
    workbook.save(result)
    return result.getvalue()


def to_csv(markdown: str) -> bytes:
    tables = markdown_tables(markdown)
    rows = tables[0][1] if tables else [["内容"], *[[line] for line in str(markdown).splitlines() if line.strip()]]
    result = io.StringIO(newline="")
    writer = csv.writer(result, lineterminator="\n")
    writer.writerows(rows)
    return b"\xef\xbb\xbf" + result.getvalue().encode("utf-8")


def render(markdown: str, fmt: str, asset_resolver=None) -> tuple[bytes, str, str]:
    fmt = str(fmt).lower().lstrip(".")
    if fmt == "docx":
        return to_docx(markdown, asset_resolver), "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if fmt == "pdf":
        return to_pdf(markdown, asset_resolver), "pdf", "application/pdf"
    if fmt == "xlsx":
        return to_xlsx(markdown), "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if fmt == "csv":
        return to_csv(markdown), "csv", "text/csv; charset=utf-8"
    raise ValueError("只支持 docx、pdf、xlsx、csv 导出。")


def project_zip(files: dict[str, bytes]) -> bytes:
    """Package validated relative paths with deflate compression."""
    result = io.BytesIO()
    with zipfile.ZipFile(result, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in sorted(files.items()):
            archive.writestr(name.replace("\\", "/"), content)
    return result.getvalue()

